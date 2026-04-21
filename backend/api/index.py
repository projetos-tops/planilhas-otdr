# backend/api/index.py
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from typing import List, Dict
import io
import re
from pathlib import Path

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.post("/api/processar-pasta")
async def processar_pasta(
    files: List[UploadFile] = File(...),
    comprimento_bobina: float = Form(...),
    operador: str = Form(...),
    nome_pasta: str = Form(...)
):
    """Processa uma pasta e retorna o Excel"""
    
    try:
        dados_fibras = []
        
        for file in files:
            content = await file.read()
            
            # Extrair dados básicos do PDF
            texto = content.decode('latin-1', errors='ignore')
            
            # Extrair declive
            declive = 0.0
            declive_match = re.search(r'Declive\s+([\d\.]+)', texto)
            if declive_match:
                declive = float(declive_match.group(1))
            
            # Extrair distância
            distancia = 0.0
            dist_match = re.search(r'(\d+\.\d+)\s*m', texto)
            if dist_match:
                distancia = float(dist_match.group(1))
            
            # Extrair número da fibra
            fibra = 0
            fibra_match = re.search(r'F[-_]?(\d+)', file.filename, re.I)
            if fibra_match:
                fibra = int(fibra_match.group(1))
            
            dados_fibras.append({
                "fibra": fibra,
                "declive": declive,
                "distancia_final": distancia,
                "data": "",
                "tipo": ""
            })
        
        # Ordenar por fibra
        dados_fibras.sort(key=lambda x: x["fibra"])
        
        # Criar Excel
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        
        # Preencher cabeçalho
        ws['A8'] = f"BOBINA {nome_pasta}"
        ws['G11'] = int(dados_fibras[0]["distancia_final"] - comprimento_bobina) if dados_fibras else 0
        ws['A67'] = operador
        
        # Preencher declives
        linha = 15
        for dado in dados_fibras:
            ws[f'E{linha}'] = dado["declive"]
            linha += 1
        
        # Salvar
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=BOBINA_{nome_pasta}.xlsx"}
        )
        
    except Exception as e:
        print(f"Erro: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/health")
async def health():
    return {"status": "ok"}